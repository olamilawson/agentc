"""Document parsing: no model involved (PRD: read client files and parse)."""

from __future__ import annotations

import io

import pytest

from personal_agent.clients import parse_document


def test_txt_and_md():
    assert parse_document(b"# brief\nobjective", "brief.md")["text"].endswith("objective")
    assert parse_document(b"plain", "notes.txt")["kind"] == "text"


def test_pdf():
    parsed = parse_document(_minimal_pdf("Hello brief"), "brief.pdf")
    assert parsed["kind"] == "pdf"
    assert parsed["pages"] == 1
    assert "Hello brief" in parsed["text"]


def test_docx():
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Q2 campaign proposal")
    doc.add_paragraph("Budget: 4.5m naira")
    doc.save(buf)
    parsed = parse_document(buf.getvalue(), "proposal.docx")
    assert parsed["kind"] == "docx"
    assert "4.5m naira" in parsed["text"]


def test_unsupported_type_rejected():
    with pytest.raises(ValueError):
        parse_document(b"MZ...", "thing.exe")


def _minimal_pdf(text: str) -> bytes:
    """A tiny valid single-page PDF with a computed xref table (no fixtures needed)."""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj".encode() + body + b"endobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer<</Size {len(objs) + 1}/Root 1 0 R>>\n"
        f"startxref\n{xref_at}\n%%EOF"
    ).encode()
    return out
